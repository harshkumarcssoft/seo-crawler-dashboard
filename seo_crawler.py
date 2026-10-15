import argparse, hashlib, ipaddress, json, math, re, socket, time
from collections import deque, defaultdict
from datetime import datetime, timezone
from urllib.parse import urlparse, urljoin, urldefrag, urlencode, parse_qsl, urlunparse, quote
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ['Address', 'Content Type', 'Status Code', 'Status', 'Indexability', 'Indexability Status', 'Title 1', 'Title 1 Length', 'Title 1 Pixel Width', 'Meta Description 1', 'Meta Description 1 Length', 'Meta Description 1 Pixel Width', 'Meta Keywords 1', 'Meta Keywords 1 Length', 'H1-1', 'H1-1 Length', 'H1-2', 'H1-2 Length', 'H2-1', 'H2-1 Length', 'H2-2', 'H2-2 Length', 'Meta Robots 1', 'X-Robots-Tag 1', 'Meta Refresh 1', 'Canonical Link Element 1', 'rel="next" 1', 'rel="prev" 1', 'HTTP rel="next" 1', 'HTTP rel="prev" 1', 'amphtml Link Element', 'Size (bytes)', 'Transferred (bytes)', 'Total Transferred (bytes)', 'CO2 (mg)', 'Carbon Rating', 'Word Count', 'Sentence Count', 'Average Words Per Sentence', 'Flesch Reading Ease Score', 'Readability', 'Text Ratio', 'Crawl Depth', 'Folder Depth', 'Link Score', 'Inlinks', 'Unique Inlinks', 'Unique JS Inlinks', '% of Total', 'Outlinks', 'Unique Outlinks', 'Unique JS Outlinks', 'External Outlinks', 'Unique External Outlinks', 'Unique External JS Outlinks', 'Closest Similarity Match', 'No. Near Duplicates', 'Spelling Errors', 'Grammar Errors', 'Hash', 'Response Time', 'Last Modified', 'Redirect URL', 'Redirect Type', 'Cookies', 'Language', 'HTTP Version', 'Extractor 1 1', 'Extractor 2 1', 'Extractor 2 2', 'Extractor 3 1', 'Extractor 3 2', 'Mobile Alternate Link', 'URL Encoded Address', 'Crawl Timestamp']
UA = "SiteAuditCrawler/1.0 (+respectful technical SEO audit)"
TIMEOUT = (8, 18)
MAX_BYTES = 4_000_000
TRACKING = {"utm_source","utm_medium","utm_campaign","utm_term","utm_content","gclid","fbclid","msclkid"}
NON_HTML = (".jpg",".jpeg",".png",".gif",".webp",".svg",".ico",".pdf",".zip",".gz",".css",".js",".mp4",".mp3",".woff",".woff2",".ttf",".eot")

def utcnow():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")

def public_host(host):
    if not host or host.lower() in {"localhost","localhost.localdomain"} or host.endswith(".local"):
        return False
    try:
        ips=[ipaddress.ip_address(host)]
    except ValueError:
        try:
            ips=[ipaddress.ip_address(x[4][0]) for x in socket.getaddrinfo(host,None,type=socket.SOCK_STREAM)]
        except Exception:
            return False
    return all(ip.is_global for ip in ips)

def host_family(host):
    h=(host or "").lower().rstrip(".")
    return h[4:] if h.startswith("www.") else h

def normalize(raw):
    try:
        raw=urldefrag(str(raw).strip())[0]
        p=urlparse(raw)
        if p.scheme.lower() not in ("http","https") or not p.hostname or not public_host(p.hostname):
            return None
        scheme=p.scheme.lower(); host=p.hostname.lower().rstrip("."); port=p.port
        if (scheme=="http" and port==80) or (scheme=="https" and port==443): port=None
        netloc=host if port is None else f"{host}:{port}"
        path=p.path or "/"
        query=urlencode([(k,v) for k,v in parse_qsl(p.query,keep_blank_values=True) if k.lower() not in TRACKING],doseq=True)
        return urlunparse((scheme,netloc,path,"",query,""))
    except Exception:
        return None

def same_site(url, family):
    try: return host_family(urlparse(url).hostname)==family
    except Exception: return False

def txt(node):
    return node.get_text(" ",strip=True) if node else ""

def meta(soup, name):
    t=soup.find("meta",attrs={"name":re.compile("^"+re.escape(name)+"$",re.I)})
    return (t.get("content") or "").strip() if t else ""

def rel_link(soup, rel, final):
    for t in soup.find_all("link",href=True):
        rs=t.get("rel",[]); rs=[rs] if isinstance(rs,str) else rs
        if rel.lower() in [str(x).lower() for x in rs]:
            return urljoin(final,t["href"])
    return ""

def pixel_width(text):
    # Approximation based on common sans-serif glyph widths; not a browser-rendered measurement.
    widths={"i":4,"l":4,"I":4,"t":5,"f":5,"r":5,"m":11,"w":10,"M":11,"W":12," ":4}
    return round(sum(widths.get(c,7) for c in (text or "")))

def flesch(text, words, sentences):
    if words == 0 or sentences == 0: return ""
    def syllables(word):
        w=re.sub(r"[^a-z]","",word.lower())
        if not w: return 0
        groups=re.findall(r"[aeiouy]+",w)
        count=len(groups)
        if w.endswith("e") and count>1 and not w.endswith(("le","ye")): count-=1
        return max(count,1)
    total=sum(syllables(w) for w in re.findall(r"\b[\w'-]+\b",text.lower()))
    score=206.835-1.015*(words/sentences)-84.6*(total/words)
    return round(score,2)

def readability(score):
    if score=="": return ""
    if score>=90: return "Very Easy"
    if score>=80: return "Easy"
    if score>=70: return "Fairly Easy"
    if score>=60: return "Standard"
    if score>=50: return "Fairly Difficult"
    if score>=30: return "Difficult"
    return "Very Difficult"

def sitemap_urls(session, origin, family, limit):
    queue=deque([urljoin(origin,"/sitemap.xml"),urljoin(origin,"/sitemap_index.xml")])
    seen=set(); found=[]
    while queue and len(found)<limit:
        sm=queue.popleft(); n=normalize(sm)
        if not n or n in seen or not same_site(n,family): continue
        seen.add(n)
        try:
            r=session.get(sm,timeout=TIMEOUT,allow_redirects=True,headers={"User-Agent":UA})
            if r.status_code>=400: continue
            soup=BeautifulSoup(r.content[:MAX_BYTES],"xml")
            if soup.find("sitemapindex"):
                queue.extend([x.get_text(strip=True) for x in soup.select("sitemap > loc")])
            else:
                for loc in soup.select("url > loc"):
                    u=normalize(loc.get_text(strip=True))
                    if u and same_site(u,family): found.append(u)
                    if len(found)>=limit: break
        except requests.RequestException: pass
    return found

def extract_page(url, response, final_url, chain, body, depth, all_rows, inlinks):
    ctype=response.headers.get("Content-Type","")
    soup=BeautifulSoup(body,"lxml") if "html" in ctype.lower() else None
    row={h:"" for h in HEADERS}
    row.update({"Address":url,"Content Type":ctype,"Status Code":response.status_code,
                "Status": "Success" if 200<=response.status_code<300 else ("Redirect" if 300<=response.status_code<400 else ("Client Error" if response.status_code<500 else "Server Error")),
                "Size (bytes)":len(body),"Transferred (bytes)":len(body),
                "Total Transferred (bytes)":len(body),"Response Time":round(response.elapsed.total_seconds(),3),
                "Last Modified":response.headers.get("Last-Modified",""),"Redirect URL":final_url if final_url!=url else "",
                "Redirect Type": " -> ".join(chain) if len(chain)>1 else "",
                "Cookies":"; ".join(response.headers.get("Set-Cookie","").splitlines()),
                "HTTP Version":getattr(response.raw,"version",""),
                "URL Encoded Address":quote(url,safe=":/?&=%#[]@!$'()*+,;"),
                "Crawl Timestamp":utcnow(),"Crawl Depth":depth,
                "Folder Depth":len([x for x in urlparse(url).path.split("/") if x]),
                "Inlinks":inlinks.get(url,0),"Unique Inlinks":inlinks.get(url,0),
                "Unique JS Inlinks":0,"% of Total":"","Link Score":"",
                "Outlinks":0,"Unique Outlinks":0,"Unique JS Outlinks":0,
                "External Outlinks":0,"Unique External Outlinks":0,"Unique External JS Outlinks":0,
                "Closest Similarity Match":"","No. Near Duplicates":0,
                "Spelling Errors":"","Grammar Errors":"","Hash":hashlib.sha256(body).hexdigest()})
    links=[]
    if not soup:
        row["Indexability"]="Non-HTML"
        row["Indexability Status"]="Not HTML"
        return row,links
    title_tag=soup.find("title"); title=txt(title_tag)
    desc=meta(soup,"description"); keywords=meta(soup,"keywords")
    h1=[txt(x) for x in soup.find_all("h1")]; h2=[txt(x) for x in soup.find_all("h2")]
    robots=meta(soup,"robots")
    xrobots=response.headers.get("X-Robots-Tag","")
    noindex="noindex" in (robots+" "+xrobots).lower()
    canonical=rel_link(soup,"canonical",final_url)
    next_link=rel_link(soup,"next",final_url); prev_link=rel_link(soup,"prev",final_url)
    refresh=next((t.get("content","") for t in soup.find_all("meta",attrs={"http-equiv":re.compile("^refresh$",re.I)})),"")
    text=txt(soup.body or soup)
    words=re.findall(r"\b[\w'-]+\b",text)
    sentences=[s for s in re.split(r"[.!?]+",text) if s.strip()]
    avg=round(len(words)/len(sentences),2) if sentences else ""
    score=flesch(text,len(words),len(sentences))
    text_ratio=round(len(text)/max(len(body),1),4)
    lang=(soup.html.get("lang","") if soup.html else "")
    images=soup.find_all("img")
    link_entries=[]
    for a in soup.find_all("a",href=True):
        target=normalize(urljoin(final_url,a["href"]))
        if target:
            internal=host_family(urlparse(target).hostname)==host_family(urlparse(url).hostname)
            link_entries.append((target,internal))
            links.append(target)
    internal=[u for u,i in link_entries if i]; external=[u for u,i in link_entries if not i]
    jsonld=[]
    for tag in soup.find_all("script",attrs={"type":re.compile("application/ld\\+json",re.I)}):
        try:
            obj=json.loads(tag.string or tag.get_text())
            for item in (obj if isinstance(obj,list) else [obj]):
                if isinstance(item,dict) and item.get("@type"):
                    jsonld.extend(item["@type"] if isinstance(item["@type"],list) else [item["@type"]])
        except Exception: pass
    alternates=[]
    for t in soup.find_all("link",href=True):
        rs=t.get("rel",[]); rs=[rs] if isinstance(rs,str) else rs
        if "alternate" in rs and t.get("hreflang"): alternates.append(t.get("hreflang")+"="+urljoin(final_url,t["href"]))
    h1s=h1+[""]*2; h2s=h2+[""]*2
    row.update({
        "Indexability":"Non-Indexable" if noindex else "Indexable (directive check only)",
        "Indexability Status":"Blocked by noindex" if noindex else "No noindex detected",
        "Title 1":title,"Title 1 Length":len(title),"Title 1 Pixel Width":pixel_width(title),
        "Meta Description 1":desc,"Meta Description 1 Length":len(desc),
        "Meta Description 1 Pixel Width":pixel_width(desc),"Meta Keywords 1":keywords,
        "Meta Keywords 1 Length":len(keywords),
        "H1-1":h1s[0],"H1-1 Length":len(h1s[0]),"H1-2":h1s[1],"H1-2 Length":len(h1s[1]),
        "H2-1":h2s[0],"H2-1 Length":len(h2s[0]),"H2-2":h2s[1],"H2-2 Length":len(h2s[1]),
        "Meta Robots 1":robots,"X-Robots-Tag 1":xrobots,"Meta Refresh 1":refresh,
        "Canonical Link Element 1":canonical,'rel="next" 1':next_link,'rel="prev" 1':prev_link,
        "amphtml Link Element":rel_link(soup,"amphtml",final_url),
        "Word Count":len(words),"Sentence Count":len(sentences),
        "Average Words Per Sentence":avg,"Flesch Reading Ease Score":score,
        "Readability":readability(score) if score!="" else "","Text Ratio":text_ratio,
        "Outlinks":len(link_entries),"Unique Outlinks":len(set(links)),
        "External Outlinks":len(external),"Unique External Outlinks":len(set(external)),
        "Language":lang,"Extractor 1 1":", ".join(sorted(set(map(str,jsonld)))),
        "Mobile Alternate Link":next((urljoin(final_url,t["href"]) for t in soup.find_all("link",href=True)
            if "alternate" in (t.get("rel",[]) if isinstance(t.get("rel",[]),list) else [t.get("rel")]) and t.get("media")),""),
    })
    return row,links

def write_excel(rows, out):
    wb=Workbook(); ws=wb.active; ws.title="1 - All"
    ws.append(HEADERS)
    for row in rows: ws.append([row.get(h,"") for h in HEADERS])
    fill=PatternFill("solid",fgColor="17365D")
    for c in ws[1]:
        c.font=Font(color="FFFFFF",bold=True); c.fill=fill
        c.alignment=Alignment(wrap_text=True,vertical="center")
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; ws.row_dimensions[1].height=32
    for i,h in enumerate(HEADERS,1):
        width=min(max(len(h)+2,14),30)
        if h in ("Address","Canonical Link Element 1","Meta Description 1","Title 1","Redirect URL"): width=38
        ws.column_dimensions[get_column_letter(i)].width=width
    wb.save(out)

def run_crawl(start_url,max_pages,delay,output,use_sitemap=True):
    start=normalize(start_url)
    if not start: raise SystemExit("Please enter a full public URL starting with https:// or http://")
    family=host_family(urlparse(start).hostname)
    origin=f"{urlparse(start).scheme}://{urlparse(start).netloc}"
    session=requests.Session(); session.headers.update({"User-Agent":UA,"Accept":"text/html,application/xhtml+xml,*/*;q=0.8"})
    robots=RobotFileParser()
    try:
        rr=session.get(origin+"/robots.txt",timeout=TIMEOUT,allow_redirects=True)
        robots.parse(rr.text.splitlines() if rr.status_code<400 else [])
    except requests.RequestException:
        raise SystemExit("Could not fetch robots.txt. Crawl stopped.")
    queue=deque(); queued=set(); visited=set(); depths={}; rows=[]; incoming=defaultdict(int); discovered_links={}
    def enqueue(raw,depth):
        u=normalize(urljoin(start,raw))
        if u and same_site(u,family) and u not in queued and len(queued)<max_pages*20:
            if not urlparse(u).path.lower().endswith(NON_HTML):
                queued.add(u); queue.append(u); depths[u]=depth
    enqueue(start,0)
    if use_sitemap:
        for u in sitemap_urls(session,origin,family,max_pages*5): enqueue(u,0)
    last=0
    while queue and len(visited)<max_pages:
        url=queue.popleft()
        if url in visited: continue
        visited.add(url)
        if not robots.can_fetch(UA,url):
            row={h:"" for h in HEADERS}; row.update({"Address":url,"Indexability":"Blocked","Indexability Status":"Disallowed by robots.txt","Crawl Depth":depths.get(url,0),"Crawl Timestamp":utcnow()})
            rows.append(row); continue
        pause=delay-(time.monotonic()-last)
        if pause>0: time.sleep(pause)
        started=time.monotonic(); last=started
        chain=[url]; current=url; response=None
        try:
            for _ in range(10):
                response=session.get(current,timeout=TIMEOUT,allow_redirects=False,stream=True)
                if response.is_redirect or response.is_permanent_redirect:
                    loc=response.headers.get("Location"); code=response.status_code; response.close()
                    if not loc: break
                    nxt=normalize(urljoin(current,loc))
                    if not nxt: break
                    chain.append(nxt); current=nxt
                    if not same_site(current,family): break
                    continue
                break
            if response is None: raise requests.RequestException("No response")
            body=response.raw.read(MAX_BYTES,decode_content=True); response.close()
            row,links=extract_page(url,response,current,chain,body,depths.get(url,0),rows,incoming)
            row["Response Time"]=round(time.monotonic()-started,3)
            discovered_links[url]=links
            for target in links:
                if same_site(target,family):
                    incoming[target]+=1
                    enqueue(target,depths.get(url,0)+1)
            rows.append(row)
        except requests.RequestException as e:
            row={h:"" for h in HEADERS}; row.update({"Address":url,"Status":"Fetch Error","Indexability":"Unknown","Indexability Status":str(e),"Crawl Depth":depths.get(url,0),"Crawl Timestamp":utcnow()}); rows.append(row)
        if len(visited)%10==0:
            write_excel(rows,output)
            print(f"PROGRESS {len(visited)} / {max_pages} pages")
    # Finalize link counts and duplicate content flags.
    hashes=defaultdict(list); titles=defaultdict(list)
    for row in rows:
        if row.get("Hash"): hashes[row["Hash"]].append(row)
        if row.get("Title 1"): titles[str(row["Title 1"]).strip().casefold()].append(row)
    total=max(len(rows),1)
    for row in rows:
        row["Inlinks"]=incoming.get(row.get("Address",""),0)
        row["Unique Inlinks"]=row["Inlinks"]
        row["% of Total"]=round(100/total,3)
        row["Link Score"]=round(math.log1p(row["Inlinks"]),3)
        near=hashes.get(row.get("Hash",""),[])
        row["No. Near Duplicates"]=max(0,len(near)-1)
        row["Closest Similarity Match"]=next((x["Address"] for x in near if x["Address"]!=row["Address"]),"")
        row["CO2 (mg)"]=""; row["Carbon Rating"]=""; row["Unique JS Inlinks"]=0
    write_excel(rows,output)
    print(f"FINISHED {len(rows)} pages. Excel report: {output}")
    return len(rows)

def main():
    p=argparse.ArgumentParser()
    p.add_argument("url"); p.add_argument("--max-pages",type=int,default=500)
    p.add_argument("--delay",type=float,default=1.0); p.add_argument("--output",default="seo_audit_report.xlsx")
    p.add_argument("--no-sitemap",action="store_true")
    a=p.parse_args()
    if not 1<=a.max_pages<=100000: p.error("--max-pages must be between 1 and 100000")
    if a.delay<0.5: p.error("--delay must be at least 0.5 seconds")
    run_crawl(a.url,a.max_pages,a.delay,a.output,not a.no_sitemap)
if __name__=="__main__": main()
